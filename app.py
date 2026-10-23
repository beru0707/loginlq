import os
import sys
import time
import threading
import subprocess
from datetime import datetime

import customtkinter as ctk
from tkinter import ttk, filedialog, messagebox
import ctypes

# Khai báo AppUserModelID riêng để Windows nhận diện Icon mới trên Taskbar
try:
    ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("mycompany.ldplayer.automation.1.0")
except Exception:
    pass

# Cấu hình chủ đề giao diện
ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")


# ============================================================
# CONFIG & CONSTANTS
# ============================================================

def resource_path(name):
    """Đường dẫn file đi kèm (icon...) - chạy được cả khi đóng gói PyInstaller."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, name)


# Thư mục chứa file .exe (hoặc file .py) - nơi đặt accounts.xlsx
APP_DIR = os.path.dirname(sys.executable if getattr(sys, "frozen", False) else os.path.abspath(__file__))
EXCEL_FILE = os.path.join(APP_DIR, "accounts.xlsx")

# Vị trí các cột trong file Excel
COL_USER, COL_PASS, COL_DEV, COL_STATUS, COL_TIME = range(5)
MIN_COLS = 5


def cell_to_str(val):
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val)


def read_excel_rows(filename):
    """Đọc toàn bộ sheet đầu tiên thành list các list (nhẹ hơn pandas rất nhiều)."""
    from openpyxl import load_workbook  # import lười
    wb = load_workbook(filename, read_only=True, data_only=True)
    try:
        rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    return rows

# Tọa độ các nút trên LDPlayer (1280 x 720)
GARENA_BTN_POS = (240, 590)  # Nút màu vàng "Đăng nhập Garena"
USERNAME_POS = (640, 230)    # Ô nhập Username
PASSWORD_POS = (640, 320)    # Ô nhập Password
LOGIN_POS = (640, 478)       # Nút "Login Now"

WAIT_CLICK_GARENA = 3.0     # Chờ trang đăng nhập mở
WAIT_INPUT = 0.5            # Chờ giữa các thao tác nhập
WAIT_LOGIN = 2.0            # Chờ sau khi bấm Login


# ============================================================
# WORKER
# ============================================================

class Worker:
    def __init__(self, serial, controller):
        self.serial = serial
        self.controller = controller
        self.device = None

    def connect(self):
        try:
            import uiautomator2 as u2  # Lazy import
            self.controller.log(f"[{self.serial}] Đang kết nối...")
            self.device = u2.connect(self.serial)
            self.device.info
            self.controller.log(f"[{self.serial}] Kết nối OK")
            return True
        except Exception as e:
            self.controller.log(f"[{self.serial}] Lỗi kết nối: {e}")
            return False

    def click_garena_button(self):
        """Bấm nút 'Đăng nhập Garena' ở màn hình đầu"""
        if self.device is None:
            raise RuntimeError("Thiết bị chưa kết nối")

        self.controller.log(f"[{self.serial}] Click nút 'Đăng nhập Garena'")
        self.device.click(GARENA_BTN_POS[0], GARENA_BTN_POS[1])
        time.sleep(WAIT_CLICK_GARENA)

    def fill_account(self, username, password):
        """Nhập tài khoản & mật khẩu"""
        if self.device is None:
            raise RuntimeError("Thiết bị chưa kết nối")

        # Username
        self.device.click(USERNAME_POS[0], USERNAME_POS[1])
        time.sleep(WAIT_INPUT)
        self.device.clear_text()
        self.device.send_keys(str(username))
        time.sleep(WAIT_INPUT)

        # Password
        self.device.click(PASSWORD_POS[0], PASSWORD_POS[1])
        time.sleep(WAIT_INPUT)
        self.device.clear_text()
        self.device.send_keys(str(password))
        time.sleep(WAIT_INPUT)

    def click_login(self):
        """Bấm nút Login Now"""
        if self.device is None:
            raise RuntimeError("Thiết bị chưa kết nối")

        self.device.click(LOGIN_POS[0], LOGIN_POS[1])
        self.controller.log(f"[{self.serial}] Đã click Login Now")
        time.sleep(WAIT_LOGIN)

    def run(self, item):
        index = item["index"]
        row_id = item["id"]
        username = item["username"]
        password = item["password"]

        self.controller.log(f"[{self.serial}] ID={row_id} | Đang xử lý: {username}")

        if not self.connect():
            self.controller.update_result(index, "CONNECT_ERROR", self.serial)
            self.controller.worker_finished(self.serial)
            return

        try:
            self.click_garena_button()
            self.fill_account(username, password)
            self.controller.log(f"[{self.serial}] ID={row_id} Đã nhập xong tài khoản: {username}")
            self.click_login()

            self.controller.update_result(index, "LOGIN_CLICKED", self.serial)
            self.controller.log(f"[{self.serial}] ID={row_id} Hoàn thành: {username}")
        except Exception as e:
            self.controller.update_result(index, "ERROR", self.serial)
            self.controller.log(f"[{self.serial}] ID={row_id} Lỗi tài khoản {username}: {e}")
        finally:
            self.controller.worker_finished(self.serial)


# ============================================================
# CONTROLLER / MAIN UI
# ============================================================

class Controller(ctk.CTk):
    def __init__(self):
        super().__init__()

        if os.path.exists(resource_path("icon.ico")):
            try:
                self.iconbitmap(resource_path("icon.ico"))
            except Exception:
                pass

        self.running = False
        self.paused = False
        self.rows = None
        self.next_position = 0
        self.start_id = None

        self.selected_devices = []
        self.device_checkboxes = {}
        self.finished_workers = set()
        self.lock = threading.Lock()
        self.save_lock = threading.Lock()

        self.create_ui()
        
        # TỐI ƯU KHỞI ĐỘNG: Đưa việc quét thiết bị và đọc file Excel chạy ngầm 
        # để cửa sổ giao diện hiện lên lập tức mà không bị đơ lúc mở app.
        threading.Thread(target=self.init_background_tasks, daemon=True).start()

    def init_background_tasks(self):
        # Nạp sẵn thư viện nặng khi người dùng đã thấy giao diện,
        # để lần bấm START đầu tiên không bị khựng.
        for mod in ("uiautomator2", "openpyxl"):
            try:
                __import__(mod)
            except Exception:
                pass
        self.refresh_devices()
        if os.path.exists(EXCEL_FILE):
            self.load_and_preview_excel(EXCEL_FILE)

    def create_ui(self):
        self.title("LDPlayer Automation Controller")
        self.geometry("920x760")
        self.minsize(850, 650)

        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(2, weight=3) # Table
        self.grid_rowconfigure(4, weight=2) # Log

        # ====================================================
        # 1. KHU VỰC CẤU HÌNH & CHỌN THIẾT BỊ
        # ====================================================
        config_frame = ctk.CTkFrame(self, corner_radius=10)
        config_frame.grid(row=0, column=0, padx=15, pady=(15, 10), sticky="ew")
        config_frame.grid_columnconfigure(1, weight=1)

        # Line 1: Chọn file Excel
        ctk.CTkLabel(config_frame, text="File Excel:", font=("Arial", 12, "bold")).grid(row=0, column=0, padx=10, pady=6, sticky="w")
        self.file_var = ctk.StringVar(value=EXCEL_FILE)
        self.file_entry = ctk.CTkEntry(config_frame, textvariable=self.file_var)
        self.file_entry.grid(row=0, column=1, padx=5, pady=6, sticky="ew")
        ctk.CTkButton(config_frame, text="Chọn File", width=90, command=self.select_file).grid(row=0, column=2, padx=10, pady=6)

        # Line 2: Ô nhập ID Bắt đầu
        ctk.CTkLabel(config_frame, text="ID bắt đầu:", font=("Arial", 12, "bold")).grid(row=1, column=0, padx=10, pady=6, sticky="w")
        id_sub_frame = ctk.CTkFrame(config_frame, fg_color="transparent")
        id_sub_frame.grid(row=1, column=1, columnspan=2, sticky="w")
        
        self.start_id_var = ctk.StringVar(value="1")
        ctk.CTkEntry(id_sub_frame, textvariable=self.start_id_var, width=80).pack(side="left", padx=(5, 10))
        ctk.CTkLabel(id_sub_frame, text="Ví dụ: 10 → bắt đầu chạy từ ID 10", text_color="gray70").pack(side="left")

        # Line 3: Chọn danh sách LDPlayer
        ctk.CTkLabel(config_frame, text="Chọn LDPlayer:", font=("Arial", 12, "bold")).grid(row=2, column=0, padx=10, pady=6, sticky="nw")
        
        dev_container = ctk.CTkFrame(config_frame, fg_color="transparent")
        dev_container.grid(row=2, column=1, columnspan=2, sticky="ew", padx=5, pady=2)
        
        self.device_list_frame = ctk.CTkFrame(dev_container, fg_color="#1e1e1e", corner_radius=6)
        self.device_list_frame.pack(side="left", fill="x", expand=True)

        ctk.CTkButton(dev_container, text="Quét lại", width=80, fg_color="#3a3a3a", hover_color="#4a4a4a", command=lambda: threading.Thread(target=self.refresh_devices, daemon=True).start()).pack(side="left", padx=(10, 0))

        # ====================================================
        # 2. KHU VỰC ĐIỀU KIỂN & TRẠNG THÁI
        # ====================================================
        control_frame = ctk.CTkFrame(self, fg_color="transparent")
        control_frame.grid(row=1, column=0, padx=15, pady=5, sticky="ew")

        self.start_btn = ctk.CTkButton(control_frame, text="START", fg_color="#2eb85c", hover_color="#1b8a3e", width=100, font=("Arial", 12, "bold"), command=self.start)
        self.start_btn.pack(side="left", padx=(0, 5))

        self.pause_btn = ctk.CTkButton(control_frame, text="PAUSE", fg_color="#f9b115", hover_color="#b98108", width=100, font=("Arial", 12, "bold"), command=self.pause)
        self.pause_btn.pack(side="left", padx=5)

        self.stop_btn = ctk.CTkButton(control_frame, text="STOP", fg_color="#e55353", hover_color="#a82323", width=100, font=("Arial", 12, "bold"), command=self.stop)
        self.stop_btn.pack(side="left", padx=5)

        self.status_var = ctk.StringVar(value="READY")
        self.status_lbl = ctk.CTkLabel(control_frame, textvariable=self.status_var, font=("Arial", 13, "bold"), text_color="#39C5BB")
        self.status_lbl.pack(side="right", padx=10)

        self.progress = ctk.CTkProgressBar(self, height=8)
        self.progress.grid(row=1, column=0, padx=15, pady=(40, 0), sticky="ew")
        self.progress.set(0)

        # ====================================================
        # 3. BẢNG DỮ LIỆU
        # ====================================================
        table_frame = ctk.CTkFrame(self)
        table_frame.grid(row=2, column=0, padx=15, pady=10, sticky="nsew")
        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)

        style = ttk.Style()
        style.theme_use("default")
        style.configure("Treeview", background="#2b2b2b", foreground="white", fieldbackground="#2b2b2b", rowheight=28)
        style.configure("Treeview.Heading", background="#1e1e1e", foreground="white", relief="flat", font=("Arial", 10, "bold"))
        style.map("Treeview", background=[("selected", "#1f538d")])

        columns = ("id", "username", "device", "status", "time")
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings")

        self.tree.heading("id", text="ID")
        self.tree.heading("username", text="Username")
        self.tree.heading("device", text="LDPlayer")
        self.tree.heading("status", text="Status")
        self.tree.heading("time", text="Time")

        self.tree.column("id", width=60, anchor="center")
        self.tree.column("username", width=260)
        self.tree.column("device", width=140, anchor="center")
        self.tree.column("status", width=140, anchor="center")
        self.tree.column("time", width=150, anchor="center")

        tree_scroll = ctk.CTkScrollbar(table_frame, command=self.tree.yview)
        self.tree.configure(yscrollcommand=tree_scroll.set)

        self.tree.grid(row=0, column=0, sticky="nsew", padx=(5, 0), pady=5)
        tree_scroll.grid(row=0, column=1, sticky="ns", padx=(0, 5), pady=5)

        self.tree.tag_configure("LOGIN_CLICKED", foreground="#2eb85c")
        self.tree.tag_configure("ERROR", foreground="#e55353")
        self.tree.tag_configure("CONNECT_ERROR", foreground="#e55353")
        self.tree.tag_configure("EMPTY", foreground="#7f8c8d")

        # ====================================================
        # 4. LOG CONSOLE
        # ====================================================
        ctk.CTkLabel(self, text="Nhật ký hoạt động (Log):", font=("Arial", 11, "bold")).grid(row=3, column=0, padx=15, sticky="w")

        self.log_box = ctk.CTkTextbox(self, font=("Consolas", 12), fg_color="#1a1a1a", text_color="#00FF66", corner_radius=8)
        self.log_box.grid(row=4, column=0, padx=15, pady=(0, 15), sticky="nsew")

    # ========================================================
    # DEVICE SCAN & SELECTION LOGIC
    # ========================================================

    def scan_online_devices(self):
        devices = []

        adb_path = r"C:\LDPlayer\LDPlayer9\adb.exe" 
        if not os.path.exists(adb_path):
            adb_path = r"C:\LDPlayer\LDPlayer4.0\adb.exe"
        if not os.path.exists(adb_path):
            adb_path = "adb"

        try:
            result = subprocess.run(
                [adb_path, "devices"], capture_output=True, text=True, timeout=5,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),  # không bật cửa sổ CMD
            )
            lines = result.stdout.strip().split("\n")[1:]

            for line in lines:
                if line.strip():
                    parts = line.split()
                    if len(parts) >= 2 and parts[1] == "device":
                        devices.append(parts[0])
        except Exception as e:
            self.log(f"Lỗi quét thiết bị: {e}")

        return devices

    def refresh_devices(self):
        online_devs = self.scan_online_devices()

        def update_ui_devices():
            for widget in self.device_list_frame.winfo_children():
                widget.destroy()

            self.device_checkboxes.clear()

            if not online_devs:
                ctk.CTkLabel(self.device_list_frame, text="Không tìm thấy giả lập nào đang mở", text_color="orange").pack(padx=10, pady=5)
                return

            for idx, serial in enumerate(online_devs, start=1):
                var = ctk.BooleanVar(value=True)
                chk = ctk.CTkCheckBox(self.device_list_frame, text=f"{serial} (Tab {idx})", variable=var, font=("Arial", 11, "bold"))
                chk.pack(side="left", padx=10, pady=6)
                self.device_checkboxes[serial] = var

            self.log(f"Đã cập nhật danh sách thiết bị: {', '.join(online_devs)}")

        self.after(0, update_ui_devices)

    def get_selected_devices(self):
        return [serial for serial, var in self.device_checkboxes.items() if var.get()]

    # ========================================================
    # EXCEL & PROCESS LOGIC
    # ========================================================

    def select_file(self):
        filename = filedialog.askopenfilename(title="Chọn file Excel", filetypes=[("Excel files", "*.xlsx")])
        if filename:
            self.file_var.set(filename)
            self.load_and_preview_excel(filename)

    def is_valid_value(self, val):
        if val is None:
            return False
        val_str = str(val).strip().lower()
        if not val_str or val_str in ["nan", "none", "null", "undefined"]:
            return False
        return True

    def load_and_preview_excel(self, filename):
        if not os.path.exists(filename):
            return False

        try:
            rows = read_excel_rows(filename)
        except Exception as e:
            err = str(e)  # tách ra vì biến `e` bị xóa khi hết khối except
            self.after(0, lambda: messagebox.showerror("Excel Error", f"Lỗi đọc file Excel:\n{err}"))
            return False

        width = max((len(r) for r in rows), default=0)
        if not rows or width < 2:
            self.after(0, lambda: messagebox.showerror("Lỗi", "File Excel phải có ít nhất 2 cột (Username, Password)."))
            return False

        # Cột: username, password, device, status, time (+ các cột phụ giữ nguyên)
        width = max(width, MIN_COLS)
        for r in rows:
            r.extend([None] * (width - len(r)))
            for c in (COL_DEV, COL_STATUS, COL_TIME):
                r[c] = cell_to_str(r[c])

        with self.lock:
            self.rows = rows

        self.after(0, self.load_table)
        self.after(0, self.update_initial_progress)
        self.log(f"Đã xem trước file: {os.path.basename(filename)} ({len(rows)} dòng)")
        return True

    def update_initial_progress(self):
        if not self.rows:
            self.progress.set(0)
            return

        completed_count = sum(1 for r in self.rows if str(r[COL_STATUS]).strip())
        self.progress.set(completed_count / len(self.rows))

    def find_start_position(self):
        try:
            start_id = int(self.start_id_var.get().strip())
        except ValueError:
            messagebox.showerror("Lỗi", "ID bắt đầu phải là số.")
            return False

        if start_id > len(self.rows):
            messagebox.showerror("Lỗi", f"Không tìm thấy ID >= {start_id}.")
            return False

        self.start_id = start_id
        self.next_position = max(start_id, 1) - 1
        return True

    def row_values(self, index):
        r = self.rows[index]
        return (index + 1, cell_to_str(r[COL_USER]), r[COL_DEV], r[COL_STATUS], r[COL_TIME])

    @staticmethod
    def status_tags(status):
        status = str(status).strip()
        return (status,) if status in ("LOGIN_CLICKED", "ERROR", "CONNECT_ERROR", "EMPTY") else ()

    def load_table(self):
        self.tree.delete(*self.tree.get_children())

        for index, row in enumerate(self.rows):
            self.tree.insert(
                "",
                "end",
                iid=str(index),
                values=self.row_values(index),
                tags=self.status_tags(row[COL_STATUS]),
            )

    def get_next_batch(self, batch_size):
        batch = []
        total = len(self.rows)

        while self.next_position < total and len(batch) < batch_size:
            index = self.next_position
            row = self.rows[index]
            self.next_position += 1

            username_raw = row[COL_USER]
            password_raw = row[COL_PASS]

            if not self.is_valid_value(username_raw) or not self.is_valid_value(password_raw):
                row[COL_STATUS] = "EMPTY"
                self.update_table(index)
                self.log(f"ID={index + 1} | Dữ liệu rỗng/lỗi → Bỏ qua")
                continue

            batch.append({
                "index": index,
                "id": index + 1,
                "username": str(username_raw).strip(),
                "password": str(password_raw).strip()
            })

        return batch

    def update_table(self, index):
        def update():
            if not self.tree.exists(str(index)):
                return
            self.tree.item(str(index), values=self.row_values(index), tags=self.status_tags(self.rows[index][COL_STATUS]))

        self.after(0, update)

    def update_result(self, index, status, device):
        with self.lock:
            row = self.rows[index]
            row[COL_STATUS] = status
            row[COL_DEV] = device
            row[COL_TIME] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        self.update_table(index)
        self.save_result()

        completed = self.next_position
        total = len(self.rows)
        self.after(0, lambda: self.progress.set(completed / total if total > 0 else 0))

    def save_result(self):
        if not self.rows:
            return
        try:
            from openpyxl import Workbook  # import lười: chỉ nạp khi cần lưu
            filename = self.file_var.get()
            with self.lock:
                snapshot = [list(r) for r in self.rows]
            with self.save_lock:  # tránh nhiều worker ghi file cùng lúc
                wb = Workbook(write_only=True)
                ws = wb.create_sheet()
                for r in snapshot:
                    ws.append([None if v == "" else v for v in r])
                wb.save(filename)
        except Exception as e:
            self.log(f"Lỗi lưu file Excel: {e}")

    def start(self):
        if self.running and self.paused:
            self.paused = False
            self.status_var.set("RUNNING")
            self.log("START → chạy batch tiếp theo")
            self.start_next_batch()
            return

        if self.running:
            self.log("Tool đang chạy.")
            return

        self.selected_devices = self.get_selected_devices()
        if not self.selected_devices:
            messagebox.showwarning("Cảnh báo", "Vui lòng tích chọn ít nhất 1 tab LDPlayer để chạy!")
            return

        if self.rows is None and not self.load_and_preview_excel(self.file_var.get()):
            return

        if not self.find_start_position():
            return

        self.running = True
        self.paused = False

        self.status_var.set("RUNNING")
        self.log("================================")
        self.log(f"ID bắt đầu: {self.start_id}")
        self.log(f"Các tab được chọn: {', '.join(self.selected_devices)}")
        self.log("Bắt đầu xử lý song song...")

        self.start_next_batch()

    def start_next_batch(self):
        if not self.running or self.paused:
            return

        batch_size = len(self.selected_devices)
        batch = self.get_next_batch(batch_size)
        
        if not batch:
            self.running = False
            self.paused = False
            self.status_var.set("COMPLETED")
            self.progress.set(1.0)
            self.save_result()
            self.log("================================")
            self.log("Đã hoàn thành toàn bộ danh sách.")
            return

        self.finished_workers.clear()
        self.log("================================")
        self.log(f"Batch mới: {len(batch)} dòng | Phân công cho {len(self.selected_devices)} thiết bị")

        for position, item in enumerate(batch):
            if position >= len(self.selected_devices):
                break
            serial = self.selected_devices[position]
            worker = Worker(serial, self)
            threading.Thread(target=worker.run, args=(item,), daemon=True).start()

    def worker_finished(self, serial):
        with self.lock:
            self.finished_workers.add(serial)
            count = len(self.finished_workers)

        self.log(f"[{serial}] Worker hoàn thành ({count}/{len(self.selected_devices)})")

        if count >= len(self.selected_devices):
            self.after(0, self.batch_finished)

    def batch_finished(self):
        if not self.running:
            return

        self.paused = True
        self.status_var.set("PAUSED - Bấm START để tiếp tục")
        self.log("================================")
        self.log("Các thiết bị được chọn đã hoàn thành batch. PAUSED.")
        self.save_result()

    def pause(self):
        if not self.running:
            return
        self.paused = True
        self.status_var.set("PAUSED")
        self.log("PAUSED thủ công.")

    def stop(self):
        self.running = False
        self.paused = False
        self.status_var.set("STOPPED")
        self.log("STOP.")

    def log(self, message):
        timestamp = datetime.now().strftime("%H:%M:%S")
        text = f"[{timestamp}] {message}\n"

        def write():
            self.log_box.insert("end", text)
            self.log_box.see("end")

        self.after(0, write)


if __name__ == "__main__":
    app = Controller()
    app.mainloop()
